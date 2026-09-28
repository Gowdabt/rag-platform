from __future__ import annotations

import csv
import io
import re
from html.parser import HTMLParser

from app.config.models import ParserConfig
from app.core.context import RequestContext
from app.domain.ports.parser import DocumentParser
from app.domain.registry import register
from app.domain.types import ParsedDocument, ParsedPage


class _HTMLText(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        self.parts.append(data)


@register("parser", "builtin")
class BuiltinParser(DocumentParser):
    def __init__(self, config: ParserConfig, ctx: RequestContext):
        super().__init__(config, ctx)

    async def parse(self, content: bytes, *, mime_type: str, filename: str | None = None) -> ParsedDocument:
        name = (filename or "").lower()
        mime = (mime_type or "").lower()
        if mime.startswith("text/") or name.endswith((".txt", ".md", ".markdown", ".csv", ".html", ".htm", ".json")):
            return self._parse_text(content, mime, name)
        if mime == "application/pdf" or name.endswith(".pdf"):
            return self._parse_pdf(content)
        if name.endswith(".docx") or mime.endswith("wordprocessingml.document"):
            return self._parse_docx(content)
        if name.endswith(".xlsx") or "spreadsheet" in mime:
            return self._parse_xlsx(content)
        if name.endswith(".pptx") or "presentation" in mime:
            return self._parse_pptx(content)
        text = content.decode("utf-8", errors="replace")
        return ParsedDocument(text=text, mime_type=mime or "application/octet-stream", metadata={"filename": filename})

    def _parse_text(self, content: bytes, mime: str, name: str) -> ParsedDocument:
        text = content.decode("utf-8", errors="replace")
        if name.endswith((".html", ".htm")) or "html" in mime:
            parser = _HTMLText()
            parser.feed(text)
            text = re.sub(r"\s+", " ", " ".join(parser.parts)).strip()
        if name.endswith(".csv"):
            reader = csv.reader(io.StringIO(text))
            text = "\n".join([", ".join(row) for row in reader])
        return ParsedDocument(text=text, mime_type=mime or "text/plain", metadata={"filename": name})

    def _parse_pdf(self, content: bytes) -> ParsedDocument:
        import fitz

        doc = fitz.open(stream=content, filetype="pdf")
        pages = []
        parts = []
        for i, page in enumerate(doc, start=1):
            t = page.get_text("text")
            pages.append(ParsedPage(page_number=i, text=t))
            parts.append(t)
        meta = {
            "page_count": doc.page_count,
            "title": doc.metadata.get("title") if doc.metadata else None,
            "author": doc.metadata.get("author") if doc.metadata else None,
        }
        return ParsedDocument(text="\n\n".join(parts), mime_type="application/pdf", metadata=meta, pages=pages)

    def _parse_docx(self, content: bytes) -> ParsedDocument:
        import io

        from docx import Document as Docx

        d = Docx(io.BytesIO(content))
        text = "\n".join(p.text for p in d.paragraphs)
        return ParsedDocument(text=text, mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", metadata={})

    def _parse_xlsx(self, content: bytes) -> ParsedDocument:
        import io

        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        lines = []
        for sheet in wb.worksheets:
            lines.append(f"# {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                lines.append("\t".join("" if c is None else str(c) for c in row))
        return ParsedDocument(text="\n".join(lines), mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", metadata={})

    def _parse_pptx(self, content: bytes) -> ParsedDocument:
        import io

        from pptx import Presentation

        prs = Presentation(io.BytesIO(content))
        parts = []
        pages = []
        for i, slide in enumerate(prs.slides, start=1):
            texts = [shape.text for shape in slide.shapes if hasattr(shape, "text")]
            t = "\n".join(texts)
            pages.append(ParsedPage(page_number=i, text=t))
            parts.append(t)
        return ParsedDocument(text="\n\n".join(parts), mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation", metadata={"page_count": len(pages)}, pages=pages)
